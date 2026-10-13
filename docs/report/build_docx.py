"""
Builds the editable Word version of the report: docs/report/Restaurant-Rush-Report.docx

Called from build.py after the PDF is printed, with the same template and the same captured
terminal output, so the two versions always say the same thing.

  * terminal captures and the UML diagrams are rendered to PNG with headless Chrome
  * pandoc converts the HTML to .docx
  * python-docx then tidies the result: fonts, table borders, a contents field, page numbers,
    page breaks before each section, and A3 landscape pages for the diagrams (as in the PDF)
"""
import copy
import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageChops

HERE = Path(__file__).resolve().parent
BUILD = HERE / "build"
IMG = BUILD / "img"
DIAGRAMS = HERE.parent / "diagrams"
OUT = HERE / "Restaurant-Rush-Report.docx"

FONT, MONO = "Calibri", "Consolas"
INK, ACCENT, MUTED = RGBColor(0x1F, 0x29, 0x37), RGBColor(0xB6, 0x52, 0x2B), RGBColor(0x6B, 0x72, 0x80)
WIDE = {"CLASS_DIAGRAM": ("01-class-diagram.svg", "Class diagram", 24.0),
        "SEQ_1": ("02-sequence-part-1-arrival-seating.svg", "Sequence diagram part 1", 24.0),
        "SEQ_2": ("03-sequence-part-2-order-pricing.svg", "Sequence diagram part 2", 24.0),
        "SEQ_3": ("04-sequence-part-3-ready.svg", "Sequence diagram part 3", 24.0),
        "SEQ_4": ("05-sequence-part-4-served.svg", "Sequence diagram part 4", 24.0),
        "SEQ_5": ("06-sequence-part-5-eating-payment.svg", "Sequence diagram part 5", 24.0),
        "SEQ_6": ("07-sequence-part-6-result.svg", "Sequence diagram part 6", 12.0)}
A3_W, A3_H = Cm(42.0), Cm(29.7)


def chrome_shot(chrome, page, out, width, height, scale):
    subprocess.run([chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--virtual-time-budget=3000",
                    f"--force-device-scale-factor={scale}", f"--window-size={width},{height}",
                    f"--screenshot={out}", page.as_uri()], check=True, capture_output=True)
    img = Image.open(out).convert("RGB")
    box = ImageChops.difference(img, Image.new("RGB", img.size, (255, 255, 255))).getbbox()
    if box:
        pad = int(4 * scale)
        img = img.crop((max(box[0] - pad, 0), max(box[1] - pad, 0),
                        min(box[2] + pad, img.width), min(box[3] + pad, img.height)))
    img.save(out, optimize=True)
    return img.size


def render_terminal(chrome, key, figure_html):
    page = IMG / f"{key}.html"
    lines = figure_html.count("\n") + 8
    page.write_text('<!doctype html><meta charset="utf-8"><link rel="stylesheet" href="../report.css">'
                    '<style>body{margin:0;background:#fff;width:700px}figure.term{margin:0;box-shadow:none}'
                    f'</style><body>{figure_html}</body>', encoding="utf-8")
    out = IMG / f"{key}.png"
    chrome_shot(chrome, page, out, 700, lines * 14 + 120, 2.5)
    return out


def render_svg(chrome, key, svg_name):
    svg = (DIAGRAMS / svg_name).read_text()
    vb = [float(v) for v in re.search(r'viewBox="([^"]+)"', svg).group(1).replace(",", " ").split()]
    w = 1800
    h = int(w * vb[3] / vb[2]) + 20
    page = IMG / f"{key}.html"
    page.write_text(f'<!doctype html><body style="margin:0;background:#fff">'
                    f'<img src="{(DIAGRAMS / svg_name).as_uri()}" style="width:{w}px;display:block"></body>')
    out = IMG / f"{key}.png"
    size = chrome_shot(chrome, page, out, w, h, 2)
    return out, size


# ---------------------------------------------------------------------------------------- docx helpers

def set_font(style, name=FONT, size=None, bold=None, color=None, italic=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def style_document(doc):
    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in styles]:
            set_font(styles[name], size=10.5, color=INK)
            styles[name].paragraph_format.space_after = Pt(5)
            styles[name].paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 30, INK), ("Heading 1", 17, INK), ("Heading 2", 12.5, INK),
                              ("Heading 3", 10.5, ACCENT)):
        set_font(styles[name], size=size, bold=True, color=color, italic=False)
        styles[name].paragraph_format.space_before = Pt(14 if name != "Heading 3" else 8)
        styles[name].paragraph_format.space_after = Pt(4)
        styles[name].paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Source Code", "Verbatim Char"):
        if name in [s.name for s in styles]:
            set_font(styles[name], name=MONO, size=8.5)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def style_tables(doc):
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for edge, size, color in (("top", 8, "9CA3AF"), ("bottom", 8, "9CA3AF"), ("insideH", 4, "D1D5DB"),
                                  ("left", 0, "FFFFFF"), ("right", 0, "FFFFFF"), ("insideV", 0, "FFFFFF")):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single" if size else "nil")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            borders.append(el)
        tblpr.append(borders)
        width = tblpr.find(qn("w:tblW"))
        if width is None:
            width = OxmlElement("w:tblW")
            tblpr.append(width)
        width.set(qn("w:type"), "pct")
        width.set(qn("w:w"), "5000")
        header = table.rows[0]._tr.find(qn("w:trPr")) is not None and \
            table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
        for r, row in enumerate(table.rows):
            for cell in row.cells:
                if r == 0 and header:
                    set_cell_shading(cell, "F3F4F6")
                for par in cell.paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    for run in par.runs:
                        run.font.size = Pt(9)
                        if r == 0 and header:
                            run.font.bold = True


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    lines = placeholder if isinstance(placeholder, list) else [placeholder]
    for i, line in enumerate(lines):
        result = paragraph.add_run(line)
        if i < len(lines) - 1:
            result.add_break()
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def insert_paragraph_after(par, text="", style=None):
    new = OxmlElement("w:p")
    par._p.addnext(new)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new, par._parent)
    if text:
        p.add_run(text)
    if style:
        p.style = style
    return p


def add_contents(doc):
    """A contents page after the cover (Word fills it in when the document is opened)."""
    first_heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    title = first_heading.insert_paragraph_before("Contents")
    title.style = doc.styles["Heading 1"]
    toc = insert_paragraph_after(title)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1" and p.text[:1].isdigit()]
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', headings)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    par.text = ""
    run = par.add_run("Restaurant Rush · SCMA Group 3      Page ")
    run.font.size, run.font.color.rgb = Pt(8.5), MUTED
    add_field(par, "PAGE", "1")
    for r in par.runs[1:]:
        r.font.size, r.font.color.rgb = Pt(8.5), MUTED


def section_break(par, landscape, template_sectpr):
    """End a section at `par`. The new sectPr describes the section that ends here."""
    sectpr = copy.deepcopy(template_sectpr)
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg"):
        for el in sectpr.findall(qn(tag)):
            sectpr.remove(el)
    size = sectpr.find(qn("w:pgSz"))
    margins = sectpr.find(qn("w:pgMar"))
    if landscape:
        size.set(qn("w:w"), str(A3_W.twips))
        size.set(qn("w:h"), str(A3_H.twips))
        size.set(qn("w:orient"), "landscape")
        for side in ("w:top", "w:bottom", "w:left", "w:right"):
            margins.set(qn(side), str(Cm(1.2).twips))
    ppr = par._p.get_or_add_pPr()
    ppr.append(sectpr)


def wide_sections(doc):
    """Put each diagram page (heading, caption, image) in an A3 landscape section, like the PDF."""
    body_sectpr = doc.sections[-1]._sectPr
    pars = doc.paragraphs
    marks = [i for i, p in enumerate(pars) if p._p.xpath(".//wp:docPr[starts-with(@descr, 'WIDE:')]")]
    groups = []
    for i in marks:
        start = i
        while start > 0 and pars[start].style.name != "Heading 1":
            start -= 1
        if groups and groups[-1][1] >= start - 2:          # consecutive diagram pages share one section
            groups[-1][1] = i
        else:
            groups.append([start, i])
    for start, end in groups:
        heading, image = pars[start], pars[end]
        after = pars[end + 1] if end + 1 < len(pars) else None
        brk = heading.insert_paragraph_before("")
        section_break(brk, False, body_sectpr)              # ends the portrait section before the diagrams
        section_break(image, True, body_sectpr)             # ends the landscape section after them
        heading.paragraph_format.page_break_before = False
        if after is not None and after.style.name == "Heading 1":
            after.paragraph_format.page_break_before = False


def size_pictures(doc, widths):
    """Pandoc shrinks every picture to the A4 text width; give each one the width it was rendered for and
    replace the internal marker in its alt text with the caption."""
    for shape in doc.inline_shapes:
        pr = shape._inline.docPr
        key = pr.get("descr", "")
        if key not in widths:
            continue
        width_cm, caption = widths[key]
        ratio = shape.height / shape.width
        shape.width = Cm(width_cm)
        shape.height = int(shape.width * ratio)
        pr.set("descr", caption)
        pr.set("title", caption)


def build_docx(template, shots, chrome, pandoc="pandoc"):
    IMG.mkdir(parents=True, exist_ok=True)
    filled = template
    widths = {}
    for key, value in shots.items():
        if key.startswith("TERM_"):
            caption = re.search(r"<em>(.*?)</em>", value)
            png = render_terminal(chrome, key, value)
            widths[key] = (16.5, html.unescape(caption.group(1)) if caption else "Terminal output")
            value = f'<p><img src="img/{png.name}" alt="{key}" style="width:16.5cm"></p>'
        elif key in WIDE:
            svg_name, alt, max_h = WIDE[key]
            png, (w, h) = render_svg(chrome, key, svg_name)
            height = max_h
            width = height * w / h
            if width > 38.5:
                width, height = 38.5, 38.5 * h / w
            widths[f"WIDE:{alt}"] = (width, alt)
            value = f'<p><img src="img/{png.name}" alt="WIDE:{alt}" style="width:{width:.1f}cm"></p>'
        filled = filled.replace("{{" + key + "}}", value)
    filled = re.sub(r"\{\{PAGE_\d+\}\}", "", filled)
    filled = re.sub(r"<title>.*?</title>", "", filled)
    filled = re.sub(r'<section class="toc">.*?</section>', "", filled, flags=re.S)
    filled = re.sub(r'<div class="term-art">.*?</div>', "", filled, flags=re.S)
    filled = re.sub(r'<span class="num">(.*?)</span>', r"\1 ", filled)
    page = BUILD / "report-docx.html"
    page.write_text(filled, encoding="utf-8")
    subprocess.run([pandoc, str(page), "-f", "html", "-t", "docx", "--shift-heading-level-by=-1",
                    f"--resource-path={BUILD}", "-o", str(OUT)], check=True)

    doc = Document(str(OUT))
    body = doc.sections[-1]                     # A4 portrait, 2 cm margins (pandoc leaves these unset)
    body.orientation = WD_ORIENT.PORTRAIT
    body.page_width, body.page_height = Cm(21.0), Cm(29.7)
    body.left_margin = body.right_margin = body.top_margin = body.bottom_margin = Cm(2.0)
    style_document(doc)
    style_tables(doc)
    # cover: course line, title, subtitle, then the details table (Contents starts on a new page)
    cover = doc.paragraphs[:3]
    cover[0].paragraph_format.space_before = Pt(110)
    for run in cover[0].runs:
        run.font.color.rgb, run.font.size = MUTED, Pt(12)
    cover[1].style = doc.styles["Title"]
    cover[1].paragraph_format.space_before = Pt(18)
    for run in cover[2].runs:
        run.font.size, run.font.color.rgb = Pt(13), RGBColor(0x37, 0x41, 0x51)
    cover[2].paragraph_format.space_after = Pt(28)
    add_contents(doc)
    wide_sections(doc)
    size_pictures(doc, widths)
    for section in doc.sections:
        add_footer(section)
    doc.core_properties.title = "Restaurant Rush: Group Assignment Report"
    doc.core_properties.author = "SCMA Group 3"
    doc.save(str(OUT))
    return OUT
